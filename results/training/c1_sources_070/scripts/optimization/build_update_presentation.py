#!/usr/bin/env python3
"""Build an editable optimization update from a consistent saved-result snapshot.

Default execution also renders the actual PPTX with LibreOffice and creates PDF
previews, a contact sheet, and bounds checks. The original baseline deck is never
written. Re-run after training to refresh numbers, status, figures, and notes.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import textwrap

os.environ.setdefault("OMP_NUM_THREADS", "1")
os.environ.setdefault("OPENBLAS_NUM_THREADS", "1")

import numpy as np
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
from build_presentation import Deck, NAVY, TEAL, ORANGE, GRAY, LIGHT, WHITE
from plot_progress import discover, number, REFERENCE, COLORS, COMPONENTS, official_submissions

SOURCES = {
    "challenge": "https://fusion-equilibrium-challenge.sophelio.io/",
    "rules": "https://fusion-equilibrium-challenge.sophelio.io/rules/",
    "data": "https://huggingface.co/datasets/Sophelio/fusion-equilibrium-challenge",
    "starter": "https://github.com/Sophelio/fusion-equilibrium-challenge-starter",
    "pod": "https://arxiv.org/abs/2111.05512",
    "fourier": "https://doi.org/10.1016/j.cma.2023.116300",
    "leaderboard": REFERENCE["url"],
    "submission": "https://www.codabench.org/competitions/17739/",
}


def load(path, fallback=None):
    try:
        return json.loads(path.read_text())
    except (FileNotFoundError, json.JSONDecodeError):
        return fallback


def compact_name(row):
    names = {"deeponet": "Initial DeepONet", "fourier_deeponet": "Initial Fourier",
             "pod_legacy": "POD + Ridge", "pod_fourier_legacy": "POD + Fourier",
             "pod_legacy_fluxloss": "POD, flux loss", "pod_legacy_noridge": "POD, no Ridge",
             "pod_legacy_flux_noridge": "POD, flux / no Ridge", "pod_full_v1": "Full-data POD",
             "pod_fourier_full_v1": "Full-data Fourier", "pod_li_refined_cpu": "POD + li refinement",
             "pod_currents64_cpu": "Currents-only POD", "pod_history64_cpu": "History POD",
             "pod_scaled_v2": "Scaled-feature POD", "pod_history64_robust_v3":"Robust history POD",
             "pod_scaled_li_v2":"Scaled POD + li"}
    if row["model"] == "validation_tuned_prediction_blend":
        return "Mean/shape blend*"
    if row.get("quality_provenance"):
        p=row["quality_provenance"]
        if p.get("expert","").startswith("scaled-li"):
            secondary=p.get("secondary_expert","")
            label="Scaled/geometry" if "geometry" in secondary else "Scaled/robust" if "robust" in secondary else "Scaled/currents"
            return label+ (" gate" if p["method"]=="gate" else " 75/25")
        if p.get("expert")=="Fourier":
            return "Fourier quality gate" if p["method"]=="gate" else "Fourier gated 75/25"
        if p.get("expert")=="li-refined_POD_DeepONet":
            return "li POD quality gate" if p["method"]=="gate" else "li POD gated 75/25"
        return "POD quality gate" if p["method"]=="gate" else "POD 75/25 blend"
    return names.get(Path(row["run_directory"]).name, row["name"])


def score_text(row):
    return f"{row['metrics']['S']:.6f}" if row else "Pending"


def fresh_comparison(data):
    # A zero-activation pure gate is exactly its main expert on this fold.
    rows=[row for row in data["fresh"] if not (row.get("quality_provenance",{}).get("method")=="gate"
          and row["quality_provenance"].get("gated_frames")==0)]
    return sorted(rows,key=lambda row:row["metrics"]["S"],reverse=True)[:8]


def snapshot(output):
    warnings = []
    rows, oracle, groups, _ = discover(ROOT / "results", output, warnings)
    manifest = load(ROOT / "data/advanced_v1/manifest.json", {})
    basis = load(ROOT / "data/pod_full_v1/basis.json", {})
    full_run = load(ROOT / "results/optimization/pod_full_v1/run.json", {})
    full_history = load(ROOT / "results/optimization/pod_full_v1/history.json", [])
    legacy = [row for row in rows if row["fold_label"].startswith("Legacy")]
    fresh = [row for row in rows if row["fold_label"].startswith("Fresh") and row["n_frames"]==4068]
    # A saved score at epoch zero is an initialization check, not a trained result.
    fresh_trained = []
    for row in fresh:
        history = load(Path(row["run_directory"]) / "history.json", [])
        scored = [item for item in history if item.get("epoch", 0) > 0 and number(item.get("S"))]
        if scored or row.get("quality_provenance"):
            matching = [item["epoch"] for item in scored if abs(item["S"]-row["metrics"]["S"])<1e-7]
            row["score_epoch"] = max(matching) if matching else None
            fresh_trained.append(row)
    run_names = ("pod_full_v1","pod_fourier_full_v1","pod_li_refined_cpu","pod_scaled_v2")
    runs = {name: load(ROOT/"results/optimization"/name/"run.json", {}) for name in run_names}
    histories = {name: load(ROOT/"results/optimization"/name/"history.json", []) for name in run_names}
    submissions=official_submissions(ROOT/"results",warnings)
    submission=submissions[-1] if submissions else {"submission_id":None,"status":"Not submitted","scores":[]}
    splits = manifest.get("splits", manifest)
    counts = {name: len(splits.get(name, [])) for name in ("train", "val", "audit", "legacy_val")}
    frames = {name: sum(item["n_frames"] for item in splits.get(name, [])) for name in counts}
    data = {"created_utc": datetime.now(timezone.utc).isoformat(), "rows": rows,
            "oracle": oracle, "legacy": legacy, "fresh": fresh_trained,
            "counts": counts, "frames": frames, "basis": basis,
            "full_run": {k: v for k, v in full_run.items() if k != "feature_names"},
            "full_history": full_history, "warnings": warnings,
            "runs":{name:{k:v for k,v in run.items() if k!="feature_names"} for name,run in runs.items()},
            "histories":histories,"submission":submission,
            "submissions":submissions,"leaderboard":load(ROOT/"results/optimization/leaderboard_review/public_snapshot.json",{}),
            "native_validation":load(ROOT/"results/optimization/native_validation_v3/combined/summary.json",{}),
            "frozen_v5":load(ROOT/"results/optimization/quality_gate_scaled_robust_v5/summary.json",{}),
            "audit_score":load(ROOT/"results/optimization/audit_frozen_v5/audit_score.json",{}),
            "audit_export":load(ROOT/"results/optimization/audit_frozen_v5/export.json",{}),
            "quality_gate":load(ROOT/"results/optimization/quality_gate_v1/summary.json",{}),
            "fourier_gate":load(ROOT/"results/optimization/quality_gate_fourier_v2/summary.json",{}),
            "li_gate":load(ROOT/"results/optimization/quality_gate_li_v3/summary.json",{}),
            "li_diagnostic":load(ROOT/"results/optimization/li_error_decomposition.json",{}),
            "historical_reference": REFERENCE, "sources": SOURCES,
            "compression": load(ROOT / "results/optimization/pod_compression_curve/compression_curve.json", {})}
    if legacy:
        data["best_legacy"] = max(legacy, key=lambda row: row["metrics"]["S"])
    if fresh_trained:
        data["best_fresh"] = max(fresh_trained, key=lambda row: row["metrics"]["S"])
    return data


def make_figures(data, output):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 12, "pdf.fonttype": 42,
                         "axes.spines.top": False, "axes.spines.right": False})
    figures = output / "figures"
    figures.mkdir(parents=True, exist_ok=True)
    def save(fig, name):
        for extension in ("png", "pdf"):
            fig.savefig(figures / f"{name}.{extension}", dpi=180, bbox_inches="tight")
        plt.close(fig)

    legacy = data["legacy"]
    keep = {"deeponet","fourier_deeponet","pod_legacy","pod_legacy_noridge","pod_full_v1",
            "pod_fourier_full_v1","pod_currents64_cpu","pod_li_refined_cpu"}
    shown = [row for row in legacy if Path(row["run_directory"]).name in keep or row.get("quality_provenance",{}).get("method")=="gated_blend75"]
    fig, ax = plt.subplots(figsize=(11.6, 4.2), layout="constrained")
    if shown:
        values = [row["metrics"]["S"] for row in shown]
        ax.plot(range(len(shown)), values, "o", color="#12A89D", ms=8)
        ax.step(range(len(shown)), np.maximum.accumulate(values), where="post", color="#12A89D", alpha=.4)
        for index, value in enumerate(values):
            ax.annotate(f"{value:.6f}", (index, value), xytext=(0, 10), textcoords="offset points", ha="center", fontsize=10)
        labels=[compact_name(row).replace("gated 75/25","mix + gate") for row in shown]
        ax.set_xticks(range(len(shown)), [textwrap.fill(label, 18) for label in labels], rotation=25, ha="right",fontsize=10)
        ax.set_ylim(min(values) - .025, 1.01)
    ax.axhline(REFERENCE["S"], ls=":", color="#A94442", label="Public test 0.995440 · Sep 12/13 · different data")
    ax.set_ylabel("Local composite S")
    ax.grid(axis="y", color="#DDE3E8")
    ax.legend(loc="upper left", fontsize=10)
    save(fig, "legacy_progress")

    original = next((row for row in legacy if Path(row["run_directory"]).name == "fourier_deeponet"), None)
    selected = [row for row in (original, data.get("best_legacy")) if row]
    fig, ax = plt.subplots(figsize=(10.8, 3.7), layout="constrained")
    left = np.zeros(len(selected))
    for index, (label, color) in enumerate(zip(COMPONENTS, COLORS)):
        values = np.asarray([row["deficits"][index] for row in selected])
        ax.barh(range(len(selected)), values, left=left, label=label, color=color)
        left += values
    ax.set_yticks(range(len(selected)), [compact_name(row) for row in selected])
    ax.set_xlabel("Weighted deficit from perfect score (1 − S)")
    ax.invert_yaxis()
    handles, labels = ax.get_legend_handles_labels()
    fig.legend(handles, labels, loc="outside lower center", ncol=4, fontsize=9, frameon=False)
    ax.grid(axis="x", color="#DDE3E8")
    ax.set_axisbelow(True)
    save(fig, "legacy_deficits")

    history = data["full_history"]
    fig, axes = plt.subplots(1, 2, figsize=(11.8, 3.7), layout="constrained")
    for key, label, style, color in (("train_loss", "Training", "--", "#91A8B0"),
                                     ("val_loss", "Fresh validation", "-", "#287D8E")):
        values = [(row["epoch"], row[key]) for row in history if row.get("epoch", 0) > 0 and number(row.get(key)) and row[key] > 0]
        if values:
            x, y = zip(*values)
            axes[0].plot(x, y, style, label=label, color=color)
    axes[0].set(yscale="log", xlabel="Epoch", ylabel="Recorded loss")
    axes[0].legend(fontsize=10)
    values = [(row["epoch"], row["S"]) for row in history if row.get("epoch", 0) > 0 and number(row.get("S"))]
    if values:
        x, y = zip(*values)
        axes[1].plot(x, y, "o-", color="#12A89D")
        axes[1].set_ylim(max(0, min(y) - .02), 1.005)
    else:
        axes[1].text(.5, .5, "Awaiting a scored training epoch", ha="center", transform=axes[1].transAxes)
    axes[1].axhline(REFERENCE["S"], ls=":", color="#A94442")
    axes[1].set(xlabel="Epoch", ylabel="Fresh-fold local S")
    for ax in axes:
        ax.grid(color="#DDE3E8")
    save(fig, "full_training")

    fig, axes=plt.subplots(1,2,figsize=(11.8,3.7),layout="constrained")
    history=data["histories"].get("pod_li_refined_cpu",[])
    for key,label,color in (("train_loss","Training loss","#91A8B0"),("val_loss","Validation loss","#287D8E")):
        values=[(row["epoch"],row[key]) for row in history if row.get("epoch",0)>0 and number(row.get(key)) and row[key]>0]
        if values:
            x,y=zip(*values);axes[0].plot(x,y,label=label,color=color)
    axes[0].set(xlabel="Refinement epoch",ylabel="Recorded loss",yscale="log")
    if axes[0].lines:axes[0].legend(fontsize=9)
    values=[(row["epoch"],row["S"]) for row in history if number(row.get("S"))]
    if values:
        x,y=zip(*values);axes[1].plot(x,y,"o-",color="#12A89D",label="POD + li loss")
    for row in data["fresh"]:
        if Path(row["run_directory"]).name=="pod_full_v1":
            axes[1].axhline(row["metrics"]["S"],color="#91A8B0",ls="--",label="Parent full-data POD")
    axes[1].set(xlabel="Refinement epoch",ylabel="Fresh-fold local S")
    axes[1].legend(fontsize=9)
    for ax in axes:ax.grid(color="#DDE3E8")
    save(fig,"li_refinement")

    fig,axes=plt.subplots(1,2,figsize=(11.8,3.8),layout="constrained")
    for ax,(split,title) in zip(axes,(("val","Fresh: 128 shots / 4,068 frames"),("legacy_val","Legacy: 32 shots / 1,024 frames"))):
        result=data["quality_gate"].get("splits",{}).get(split,{})
        metrics=result.get("scores",{})
        names=("full","currents","gate") if data["li_gate"] else ("full","currents","gate","blend75")
        values=[metrics.get(name,{}).get("S",np.nan) for name in names]
        labels=["Full POD","Currents","POD gate","POD 75/25"][:len(names)]
        extra=data["fourier_gate"].get("splits",{}).get(split,{}).get("scores",{})
        extra_methods=(("gated_blend75","F. gated 75/25"),) if data["li_gate"] else (("gate","Fourier gate"),("gated_blend75","F. gated 75/25"))
        for key,label in extra_methods:
            if key in extra:values.append(extra[key]["S"]);labels.append(label)
        li_extra=data["li_gate"].get("splits",{}).get(split,{}).get("scores",{})
        for key,label in (("gate","li POD gate"),("gated_blend75","li gated 75/25")):
            if key in li_extra:values.append(li_extra[key]["S"]);labels.append(label)
        if data["frozen_v5"]:
            frozen=data["frozen_v5"].get("splits",{}).get(split,{}).get("scores",{})
            fold=data["fresh"] if split=="val" else data["legacy"]
            by_name={Path(row["run_directory"]).name:row for row in fold}
            values=[by_name[name]["metrics"]["S"] for name in ("pod_full_v1","pod_currents64_cpu","pod_history64_robust_v3","pod_scaled_li_v2")]
            values += [frozen["gate"]["S"],frozen["gated_blend75"]["S"]]
            labels=["Full POD","Currents","Robust history","Scaled + li","v5 gate","v5 75/25"]
        ax.bar(range(len(values)),values,color=["#91A8B0","#287D8E","#12A89D","#E79938","#6F60A7","#CF6272"][:len(values)])
        for i,value in enumerate(values):
            if np.isfinite(value):ax.text(i,value+.00055,f"{value:.6f}",ha="center",fontsize=8)
        ax.set(xticks=range(len(values)),xticklabels=[textwrap.fill(label,10) for label in labels],ylim=(.978,1.0),ylabel="Local composite S",title=title)
        ax.tick_params(axis="x",labelsize=9)
        coverage=result.get("coverage",{})
        ax.text(.04,.06,f"Gate active: {coverage.get('gated_frames','?')}/{coverage.get('frames','?')} frames",transform=ax.transAxes,fontsize=11)
        ax.grid(axis="y",color="#DDE3E8");ax.set_axisbelow(True)
    save(fig,"quality_gate")

    fresh=fresh_comparison(data)
    fig,ax=plt.subplots(figsize=(6.8,4.0),layout="constrained")
    left=np.zeros(len(fresh))
    for index,(label,color) in enumerate(zip(COMPONENTS,COLORS)):
        values=np.asarray([row["deficits"][index] for row in fresh])
        ax.barh(range(len(fresh)),values,left=left,label=label,color=color);left+=values
    ax.set_yticks(range(len(fresh)),[compact_name(row) for row in fresh],fontsize=10)
    ax.set_xlabel("Deficit from 1 (smaller is better)",fontsize=10)
    ax.invert_yaxis();ax.grid(axis="x",color="#DDE3E8");ax.set_axisbelow(True)
    handles,labels=ax.get_legend_handles_labels()
    fig.legend(handles,labels,loc="outside lower center",ncol=2,fontsize=9,frameon=False)
    save(fig,"fresh_deficits")
    return figures


class UpdateDeck(Deck):
    def slide(self, title, kicker="OPTIMIZATION UPDATE", source="Saved local experiment artifacts", note=""):
        slide = super().slide(title, kicker, source, note)
        for shape in slide.shapes:
            if shape.has_text_frame and "INITIAL DIII-D BASELINES" in shape.text:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.text = run.text.replace("INITIAL DIII-D BASELINES", "OPTIMIZATION UPDATE")
        return slide

    def export(self, output):
        self.prs.core_properties.title = "Fusion equilibrium: optimization update"
        self.prs.core_properties.subject = "Data expansion, POD and Fourier operators, validation progress, and next experiments"
        self.prs.core_properties.author = "Fusion project"
        self.prs.save(output / "fusion_optimization_update.pptx")
        (output / "slide_content.json").write_text(json.dumps(self.notes, indent=2) + "\n")
        lines = ["# Optimization-update speaker notes", ""]
        for slide in self.notes:
            lines += [f"## {slide['slide']:02d}. {slide['title']}", "", slide["notes"], "", f"Source: {slide['source']}", ""]
        (output / "speaker_notes.md").write_text("\n".join(lines))


def build(data, output, figures):
    d = UpdateDeck()
    counts, frames, basis = data["counts"], data["frames"], data["basis"]
    legacy_best, fresh_best = data.get("best_legacy"), data.get("best_fresh")
    n_features, active, rank = basis.get("input_features", 1138), basis.get("active_features", 600), basis.get("rank", 128)
    history = data["full_history"]
    latest_epoch = max((item.get("epoch", 0) for item in history), default=0)
    run = data["full_run"]
    planned = run.get("arguments", {}).get("epochs", 600)
    status = "Completed" if run.get("status") == "completed" else "In progress"
    by_name = {Path(row["run_directory"]).name: row for row in data["legacy"]}
    fresh_by_name={Path(row["run_directory"]).name:row for row in data["fresh"] if not row.get("quality_provenance")}
    li_row=fresh_by_name.get("pod_li_refined_cpu")
    li_run=data["runs"].get("pod_li_refined_cpu",{})
    li_history=data["histories"].get("pod_li_refined_cpu",[])
    li_latest=max((row.get("epoch",0) for row in li_history),default=0)
    li_status="Completed" if li_run.get("status")=="completed" else "In progress"
    submission=data["submission"]
    submission_text=f"Submission {submission['submission_id']}: {submission['status']}" if submission.get("submission_id") else "Official submission pending"
    official_values={row.get("column_key"):row.get("score") for row in submission.get("scores",[]) if isinstance(row,dict)}
    official_score_text=f"{float(official_values['d3d_S']):.6f}" if "d3d_S" in official_values else "Pending"
    official_note=f"The official DIII-D composite S is {official_score_text}. Local validation scores remain separate from official test results."
    snapshot_note = f"This deck is a consistent saved-artifact snapshot created {data['created_utc']}. Training may have advanced since this build. Re-run the builder to refresh."

    s = d.slide("Toward a top fusion-equilibrium score", "RESEARCH UPDATE / DIII-D", "Challenge website; public leaderboard API; local score files; official submission status", snapshot_note + " The objective is an official first-place Challenge-1 result while approaching composite score 1.0. The public reference is a historical observation, not a claim about the current leaderboard. Local validation and public-test scores use different shots. " + official_note + " No leaderboard rank is inferred from local scores.")
    d.text(s,.7,1.65,7.8,1.35,"Improve DIII-D reconstruction\nand build zero-shot MAST transfer",29,NAVY,True)
    d.text(s,.7,3.2,7.6,1.1,f"{counts['train']:,} training shots · {frames['train']:,} frames\nDeepONet → POD → Fourier refinement",22,GRAY)
    d.card(s,9.0,1.65,3.6,2.05,"HISTORICAL PUBLIC TEST","0.995440\nObserved Sep 12/13, 2026",ORANGE,26)
    hero_row = fresh_best or legacy_best
    hero_title = "FRESH VALIDATION" if fresh_best else "MATCHED LEGACY FOLD"
    hero_shots = counts['val'] if fresh_best else 32
    d.card(s,9.0,3.95,3.6,2.0,hero_title,f"{score_text(hero_row)}\n{hero_shots} held-out shots",TEAL,26)
    d.text(s,.7,5.45,7.8,.72,f"Goal: approach S = 1 and verify rank through Codabench.\n{submission_text}. Official score: {official_score_text.lower()}.",18)

    s = d.slide("Predict the field and the scalars that describe the plasma", "GOAL + EVALUATION", "Official challenge/starter scorer v3.2.0; initial score files", "Inputs are permitted actuator currents, Thomson temperature/density, physical geometry, and requested EFIT times/grid. Targets are native 65×65 poloidal flux plus q95 and betaN. Score weights are from the unmodified scorer. All shape terms come from the submitted flux. The initial Fourier model's consistency deficit was 0.0962898 out of a total deficit 0.127948, approximately 75%. This motivates shape-sensitive training and composite-based model selection.")
    d.card(s,.7,1.6,5.8,1.35,"INPUT", "Currents + Thomson profiles + geometry + query time",TEAL,20)
    d.card(s,6.85,1.6,5.8,1.35,"OUTPUT", "ψ(R,Z): 65 × 65 pixels; q95 and βN: two scalars",ORANGE,20)
    for x, pct, label in ((.7,"55%","Flux R²"),(3.78,"15%","q95 / βN"),(6.86,"10%","Boundary"),(9.94,"20%","Derived shape")):
        d.rect(s,x,3.45,2.7,1.35,NAVY,rounded=True)
        d.text(s,x+.15,3.6,2.4,.5,pct,28,WHITE,True,PP_ALIGN.CENTER)
        d.text(s,x+.15,4.18,2.4,.32,label,17,WHITE,align=PP_ALIGN.CENTER)
    d.text(s,.8,5.25,11.8,.8,"Initial flux maps looked accurate, but derived shape contributed 75% of the remaining score deficit.",23,NAVY,True)

    s = d.slide("Expand the training population; preserve whole-shot separation", "DATA + VALIDATION", "data/advanced_v1/manifest.json; data/pod_full_v1/basis.json", f"The downloaded release has 7,041 labeled DIII-D shots, 874 DIII-D public-test shots, and 1,206 MAST public-test shots. The new training split contains {counts['train']} shots and {frames['train']} native sampled frames; no target interpolation was performed. Fresh development and reserved audit shots were selected outside all 288 shots touched by initial baseline experiments. The original 32-shot validation split remains excluded from training. The audit was reserved during development and then consumed once for the frozen v5 candidate after explicit authorization. It was not used for fitting or selection. Train-only POD preprocessing records its source manifest and basis provenance.")
    rows = [["Partition", "Shots", "Cached frames", "Purpose"],
            ["Training",f"{counts['train']:,}",f"{frames['train']:,}","Fit networks + normalization + POD"],
            ["Fresh validation",str(counts['val']),f"{frames['val']:,}","Select models and hyperparameters"],
            ["Reserved audit",str(counts['audit']),f"{data['audit_score'].get('evaluation',{}).get('evaluated_frames',frames['audit']):,}","One frozen-candidate audit; all native frames"],
            ["Legacy validation",str(counts['legacy_val']),f"{frames['legacy_val']:,}","Matched comparison with first results"]]
    d.table(s,.75,1.7,[2.45,1.15,2.1,6.0],rows,row_h=.65,size=18)
    d.text(s,.8,5.35,11.7,.8,"26× more training shots than the initial 256-shot run.\nPublic tests remain separate: 874 DIII-D shots + 1,206 MAST shots.",21,NAVY,True)

    s = d.slide(f"Richer input information: {n_features:,} features → {active} active", "FEATURES + PREPROCESSING", "fusion/advanced_data.py; feature_audit.json; training-only POD normalization", f"The v1 feature vector retains the original 151 features and adds padded full Thomson profiles, physical chord/coil geometry, and current context at offsets −100, −20, +20, +100 milliseconds. Core and edge capacities are 64 and 32. Constant or near-constant columns are removed using training standard deviations only, leaving {active} active features from {n_features}. Positive offsets use available future input measurements, so this is an offline reconstruction experiment, not a strict real-time predictor. Per-shot official Ip timestamp correction remains applied. Missingness masks are inputs. EFIT targets and magnetics_dsep never enter features.")
    d.card(s,.7,1.65,3.85,3.35,"ACTUATOR CONTEXT","21 current signals\nAt query time and ±20 / ±100 ms\nAvailability masks\nCorrected Ip timestamps",TEAL,21)
    d.card(s,4.75,1.65,3.85,3.35,"KINETIC PROFILES","Full Te / ne channels\nCore capacity: 64\nEdge capacity: 32\nLog values + masks",ORANGE,21)
    d.card(s,8.8,1.65,3.85,3.35,"PHYSICAL GEOMETRY","Per-shot chord R / Z\nCoil location + extent\nChannel alignment checks\nTrain-only normalization",TEAL,21)
    d.text(s,.8,5.4,11.7,.73,"Future input samples make this an offline reconstruction model.\nNo target-derived quantities are model inputs.",20,GRAY)

    s = d.slide("POD-DeepONet: a compact branch with a fixed spatial trunk", "MODEL I/O", "fusion/advanced_models.py; fusion/pod.py; POD-DeepONet primary paper", f"The current full-data branch receives a batch shaped B×{active}. A width-256 stem feeds three residual MLP blocks. The branch outputs {rank} standardized POD coefficients and two standardized scalar values. Linear skip heads remain learnable, but the full-data run starts without Ridge fitting. A train-only mean map and POD coordinate basis reconstruct ψ on the fixed 65×65 grid; per-mode coefficient standard deviations and scalar mean/std restore physical units. The current run has {run.get('parameters',1088308):,} parameters. The POD trunk is fixed, not trained on validation maps. The optional learned coordinate trunk is disabled in this run.")
    for x, title, body in ((.65,"INPUT",f"B × {active}\nnormalized features"),(3.7,"BRANCH","256-wide stem\n3 residual MLP blocks"),(6.75,"COEFFICIENTS",f"B × {rank}\nstandardized amplitudes"),(9.8,"FLUX MAP","B × 65 × 65\nphysical ψ units")):
        d.card(s,x,2.0,2.75,1.75,title,body,TEAL,20)
    for x in (3.42,6.47,9.52):
        d.arrow(s,x,2.85,x+.22,2.85)
    d.card(s,3.7,4.18,2.75,1.25,"SCALAR HEAD","B × 2 → q95, βN",ORANGE,19)
    d.arrow(s,5.07,3.8,5.07,4.12)
    d.text(s,7.0,4.23,5.3,.72,"ψ = training mean + Σ cⱼ σⱼ φⱼ(R,Z)",23,NAVY,True)
    d.text(s,7.0,5.12,5.3,.65,"Fixed trunk: 128 spatial basis maps\nFitted on training frames only",18,GRAY)
    d.text(s,.85,6.14,11.6,.43,"v5 secondary: 708 active inputs → 64 POD coefficients; 128-wide branch + current history and 10% diagnostic dropout.",17,GRAY)

    s = d.slide("Fourier-DeepONet refines the reconstructed flux spatially", "MODEL ARCHITECTURE", "fusion/advanced_models.py; pod_fourier_full_v1 scores/run; Fourier-DeepONet primary paper", "The implemented POD/Fourier variant starts from a POD branch and its reconstructed normalized field. It concatenates that field with normalized R,Z coordinates and eight conditioning channels from the branch, lifts to 16 channels, pads 65×65 to 72×72, applies two spectral/local residual blocks with ten retained modes, crops, and projects to one residual field. The final projection starts at zero so a warm start initially preserves its parent POD prediction. The reported variant is an adaptation, not an assertion of reproducing the complete seismic Fourier-DeepONet architecture. The full-data refinement completed20epochs. Compare its retained fresh and legacy scores to the same-fold parent scores; it improves fresh validation but does not improve the legacy fold.")
    for x, title, body in ((.7,"CONDITION","POD flux + R,Z\n8 branch channels"),(3.75,"LIFT + PAD","16 channels\n65² → 72²"),(6.8,"FOURIER BLOCKS","2 spectral/local blocks\n10 retained modes"),(9.85,"RESIDUAL OUTPUT","Crop → project\nAdd correction to ψ")):
        d.card(s,x,1.95,2.65,2.0,title,body,TEAL if x<6 else ORANGE,20)
    for x in (3.4,6.45,9.5):
        d.arrow(s,x,2.95,x+.26,2.95)
    d.text(s,.85,4.62,11.5,.65,f"Fresh S: {score_text(fresh_by_name.get('pod_full_v1'))} → {score_text(fresh_by_name.get('pod_fourier_full_v1'))}",26,NAVY,True)
    d.text(s,.85,5.45,11.5,.65,f"Legacy S: {score_text(by_name.get('pod_full_v1'))} → {score_text(by_name.get('pod_fourier_full_v1'))}. 20 refinement epochs completed.",21,GRAY)

    s = d.slide("The spatial representation can express near-perfect fields", "ORACLE COMPRESSION DIAGNOSTIC", "results/optimization/pod_compression_curve; basis fitted on 256 training shots", "This is explicitly an oracle capacity experiment, not learned prediction. For each rank, validation ψ is projected onto a training-derived POD basis; q95 and betaN are copied from validation labels. All 32 legacy validation shots and 1,024 cached frames are scored with the official evaluator. Rank32 yields S0.997637, rank64 0.999294, and rank128 0.999510. Thus 64 modes already provide most of the representation accuracy of 128 on this fold. Neither an oracle score nor retained training variance determines how well an input-to-coefficient network will generalize.")
    curve = ROOT / "results/optimization/pod_compression_curve/compression_curve.png"
    if curve.exists():
        d.image(s,curve,.62,1.43,12.1,4.48)
    d.text(s,.85,6.05,11.6,.52,"Oracle coefficients use validation labels. These are not submission results.",20,ORANGE,True)

    s = d.slide("Train the field to improve internal inductance", "PHYSICS-GUIDED REFINEMENT", "pod_li_refined_cpu; li_error_decomposition.json; data/li_full_v1 references", f"The parent full-data POD has fresh internal-inductance R²0.927880. An earlier256-frame validation diagnostic found that field-profile errors, rather than contour motion, dominated li error: motion-only SSE was1.80% of total. This diagnostic uses validation truth for error analysis only. Refinement adds a differentiable fixed-training-contour li loss with weight0.02, initialized from the retained full-data POD. Training references use53737 valid frames, approximately8 per training shot, and512-point contours. No validation or audit labels are used in this new training loss. At this snapshot the run is {li_status.lower()}, latest epoch{li_latest}; the retained fresh score is {score_text(li_row)}. Exact official scoring still derives li from each predicted contour, so the training surrogate is not itself the deployed metric. " + snapshot_note)
    d.card(s,.75,1.5,3.85,1.3,"TRAINING-ONLY LI LOSS","53,737 valid sampled frames",TEAL,20)
    d.card(s,4.75,1.5,3.85,1.3,"FRESH LOCAL SCORE",score_text(li_row),ORANGE,26)
    li_r2=li_row["metrics"].get("r2_cons_each",{}).get("li") if li_row else None
    d.card(s,8.75,1.5,3.85,1.3,"DERIVED LI R²",f"0.927880 → {li_r2:.6f}" if li_r2 is not None else "Awaiting scored result",TEAL,20)
    d.image(s,figures/"li_refinement.png",.7,3.0,11.95,3.05)
    d.text(s,.85,6.18,11.5,.4,f"{li_status} · epoch {li_latest}. Fixed training contours guide the loss; the official scorer uses predicted contours.",17,GRAY)

    s = d.slide("Matched legacy validation shows measurable improvement", "LOCAL RESULTS / SAME 32 SHOTS", "Saved score files; exact shot/frame fingerprint; blend summary provenance", "All points on this plot use exactly the same 32 legacy shots and 1,024 native cached frames, with the same scorer version and fingerprint. Each point is a retained run score; saved runs are ordered by evaluation timestamp. Asterisks identify a blend whose global mean/shape/scalar weights were selected on this validation fold, not a new independent holdout. The public-test horizontal reference uses different shots and does not imply an official rank. This slide automatically updates when new matched legacy results are saved. " + snapshot_note)
    d.image(s,figures/"legacy_progress.png",.65,1.5,12.0,4.4)
    d.text(s,.85,6.02,11.6,.48,"Same raw shots, native frames and scorer. All values shown are local scores.",17,GRAY)

    s = d.slide("Diagnostic quality changes which expert works best", "INPUT-ONLY GATE / TWO SEPARATE FOLDS", "quality_gate_v1, quality_gate_fourier_v2, quality_gate_li_v3: policy.json + summary.json", "A fixed input-only rule selects the currents-only POD whenever no more than10% of real edge Thomson channels have both positive observed Te and ne. Padding is excluded using geometry and observation masks. Otherwise it uses the primary expert or fixed expert blend. The same prediction rule applies to flux and both scalars; no target values, shot identity or per-shot correction enter the rule. The threshold0.1 and75/25 weight were fixed before evaluation; no grid search was used. The gate activates on61/1024 legacy frames in2shots, and0/4068 fresh frames. Thus fresh validation does not independently test the fallback branch. The earlier256-shot ablation also improved legacy S from0.875588 to0.934889 by disabling the Ridge warm start; the linear skip remained trainable. Rare zero-edge-TS cases motivated this reliability check, without establishing a universal causal mechanism.")
    d.text(s,.85,1.48,11.7,.67,"Frozen v5: use robust history for poor edge diagnostics; otherwise 75% scaled-li + 25% robust history.",22,NAVY,True)
    d.image(s,figures/"quality_gate.png",.7,2.35,11.95,3.57)
    public_coverage=data["fourier_gate"].get("public_test_input_only_coverage",{})
    public_text=f"Public inputs only: {public_coverage.get('gated_frames',0):,}/{public_coverage.get('frames',0):,} frames trigger fallback; {public_coverage.get('gated_shots',0)}/{public_coverage.get('shots',0)} shots. No labels used." if public_coverage else "Fresh validation does not independently validate the low-quality fallback."
    d.text(s,.85,6.05,11.6,.64,f"Fresh fold: zero gate activations. {public_text}",17,ORANGE,True)
    d.notes[-1]["notes"] += " Earlier Fourier and li-refined gates used the currents expert. Frozen v5 instead uses the robust-history expert; never infer secondary identity from a legacy source-dictionary key. Each gated75/25 variant uses currents in low-quality frames and0.75primary+0.25currents otherwise. Its local fresh/legacy scores are in quality_gate_fourier_v2/summary.json and quality_gate_li_v3/summary.json. Public-test coverage counts use input features only and are not an evaluation score."
    d.slides[-1].notes_slide.notes_text_frame.text=d.notes[-1]["notes"]

    s = d.slide("Fresh 128-shot validation compares the latest candidates", "LOCAL RESULTS / SAME 4,068 FRAMES", "Retained val_score.json files; exact dataset/shot/frame/scorer fingerprint", "Each table row uses the same128 fresh validation shots and4068 native frames. Training checkpoints are selected using this development fold, which is distinct from the reserved audit. Only a retained score following a positive training epoch appears for a trained model. Fixed prediction combinations are labeled explicitly; they combine frozen model outputs using predeclared rules. The consistency component still dominates much of the remaining deficit. Legacy results are on a different fold and must not be pooled with these. Public-test reference0.995440 is historical and does not establish a local leaderboard rank. " + snapshot_note)
    table_rows=[["Candidate", "Local S", "Status"]]
    shown=fresh_comparison(data)
    for row in shown:
        name=Path(row["run_directory"]).name
        status="Fixed rule" if row.get("quality_provenance") else ("Done" if row.get("status")=="completed" else "Running")
        table_rows.append([compact_name(row),score_text(row),status])
    d.table(s,.65,1.65,[3.45,1.35,1.25],table_rows,row_h=.48,size=16)
    d.image(s,figures/"fresh_deficits.png",6.75,1.65,5.95,4.05)
    d.text(s,.8,6.16,11.7,.5,"Audit was consumed once for frozen v5, separately from selection. Local validation does not establish official rank.",17,GRAY)

    native=data["native_validation"].get("splits",{}).get("val",{}).get("scores",{}).get("gated_blend75",{})
    audit=data["audit_score"]
    s=d.slide("A single frozen candidate generalizes beyond development", "FULL NATIVE FRAMES + ONE-TIME AUDIT", "native_validation_v3/combined/summary.json; audit_frozen_v5/{protocol,events,audit_score}.json", "The frozen v5 ensemble was selected using development data, then exported and structurally validated before the audit was authorized. Full-native development evaluates the same128 fresh shots on all28447frames; it is not a new independent shot fold. The audit uses128 different reserved shots and all27631nativeframes. Its protocol was recorded at04:19:44UTC before first authorized audit preparation at04:19:47; scoring completed04:21:27. Exactly one frozen candidate was scored and no tuning, checkpoint comparison or expert comparison occurred. The audit is now consumed and must not become a development fold. Both checkpoints match the official export hashes. A preserved pre-edit training source reproduces the original hash and its inference helpers exactly match current helpers. The independent audit gate never activated, so its low-quality branch remains untested on these audit shots.")
    comparison=[["Population", "Shots", "Frames", "Local S"],
        ["Fresh, sampled", "128", "4,068", f"{data['frozen_v5'].get('splits',{}).get('val',{}).get('scores',{}).get('gated_blend75',{}).get('S',0):.6f}"],
        ["Same fresh shots, all frames", "128", "28,447", f"{native.get('S',0):.6f}"],
        ["Reserved audit, all frames", "128", "27,631", f"{audit.get('S',0):.6f}"]]
    d.table(s,.8,1.62,[5.35,1.35,2.45,2.5],comparison,row_h=.68,size=20)
    d.card(s,.8,4.7,5.78,1.3,"FROZEN AUDIT RESULT",f"Flux R² {audit.get('r2_psi',0):.6f} · consistency {audit.get('consistency',0):.6f}",TEAL,20)
    d.card(s,6.8,4.7,5.78,1.3,"IMPORTANT COVERAGE LIMIT","Zero fallback activations in the audit",ORANGE,20)
    d.text(s,.85,6.2,11.6,.43,"No audit labels were used for fitting or selection. Expanded frame coverage changes the metric population.",17,GRAY)

    board=data["leaderboard"]
    first=board.get("d3d_S_top10",[{}])[0];second=board.get("g_ratio_top10",[{}])[0]
    s=d.slide("Official progress: #9 → #4; the two challenges use different objectives", "PUBLIC LEADERBOARD / TIMESTAMPED OBSERVATIONS", "leaderboard_review/public_snapshot.json; public competition configuration; official scorer v3.2.0", f"The first official submission924272 achievedDIII-D S0.993190 and was verified#9. The second frozen-v5 submission924312 achieved0.994620 and was verified#4. Current public observations are timestamped{board.get('observed_utc')}. Challenge1 ranksDIII-D S. Challenge2 ranksG=S_MAST/S_DIII-D withDIII-D S≥0.85; G is not clipped at1. Its current leader hasG1.007046 withMAST S1.0 andDIII-D S0.993003. One entry beating the current Ch1leader0.995440 hasmaximum possibleGbelow1.004581, so one prediction pair cannot beat both present leaders simultaneously. Force_Last displays the team's latest public entry; it does not retain a different row for each objective. Pursue genuine reconstruction progress in separate model tracks, without degrading DIII-D predictions to inflate the ratio. Both existing submissions contain DIII-D only, so their G0 is an omitted-MAST result, not measured model transfer accuracy.")
    official_rows=[["Our official entry", "DIII-D S", "Observed rank"]]
    for entry in data["submissions"]:
        values={v["column_key"]:v["score"] for v in entry["scores"]}
        if "d3d_S" not in values:continue
        rank=entry.get("first_rank_observation") or {}
        official_rows.append([str(entry["submission_id"]),f"{float(values['d3d_S']):.6f}",f"#{rank['rank']}" if rank else "Unverified"])
    d.table(s,.8,1.65,[5.5,3.0,3.15],official_rows,row_h=.65,size=22)
    d.card(s,.8,4.03,5.8,1.65,"CHALLENGE 1 LEADER",f"{first.get('owner','Pending')} · S {first.get('scores',{}).get('d3d_S',0):.6f}",TEAL,23)
    d.card(s,6.8,4.03,5.8,1.65,"CHALLENGE 2 LEADER",f"{second.get('owner','Pending')} · G {second.get('scores',{}).get('g_ratio',0):.6f}",ORANGE,23)
    d.text(s,.85,6.0,11.6,.63,f"Observed {board.get('observed_utc','unknown')}. G can exceed 1; one entry cannot currently lead both columns.",18,GRAY)

    s = d.slide("Track the official submission and continue controlled improvements", "OFFICIAL STATUS + NEXT EXPERIMENTS", "Saved submission status files; public leaderboard API; Codabench competition17739", f"The displayed official entry is the exact validated frozen archive identified by its saved status file. {submission_text}; created {submission.get('created_utc')}. The safe submission status artifact is read at build time, and no private authentication data are included. {official_note} Leaderboard rank is reported only from its timestamped public API observation. Remaining work is to read the official evaluator outcome, compare metric components against local validation, continue controlled model improvements, and verify generalization before choosing a subsequent artifact. The one-time audit of frozen v5 is complete and was not used for model selection. Source URLs accompany this deck and each local result is traceable through snapshot.json and its score file.")
    d.card(s,.7,1.62,5.8,3.42,"NEXT SCIENTIFIC CHECKS","Improve models on development folds\nRetain the completed audit unchanged\nBuild a transferable MAST operator\nValidate physics before submission\nPreserve reproducible code + weights",TEAL,21)
    rank=submission.get("latest_rank_observation") or {}
    rank_text=f"#{rank['rank']} at recorded snapshot" if rank else "Unverified"
    d.card(s,6.8,1.62,5.8,3.42,"OFFICIAL SUBMISSION",f"ID: {submission.get('submission_id','Pending')}\nModel: scaled/robust POD v5\nStatus: {submission['status']}\nOfficial DIII-D S: {official_score_text}\nLeaderboard: {rank_text}",ORANGE,21)
    d.text(s,.8,5.43,11.7,.92,"Sources: challenge website + rules · Hugging Face dataset · official starter/scorer\nPOD-DeepONet: arXiv:2111.05512 · Fourier-DeepONet: DOI 10.1016/j.cma.2023.116300\nCodabench competition 17739 · full URLs and local provenance accompany this deck",14,GRAY)
    references = "\n\nSource URLs:\n" + "\n".join(f"{key}: {value}" for key,value in SOURCES.items())
    d.notes[-1]["notes"] += references
    d.slides[-1].notes_slide.notes_text_frame.text = d.notes[-1]["notes"]
    d.export(output)
    return len(d.slides)


def render_and_check(output, expected):
    subprocess.run(["bash", str(ROOT/"scripts/render_presentation.sh"),
                    str(output/"fusion_optimization_update.pptx"), str(output)], check=True)
    import pymupdf as fitz
    prs = Presentation(output/"fusion_optimization_update.pptx")
    pdf = fitz.open(output/"fusion_optimization_update.pdf")
    if len(prs.slides) != expected or len(pdf) != expected:
        raise AssertionError("PPTX and actual Office PDF page counts differ")
    previews = output/"previews"
    previews.mkdir(exist_ok=True)
    outside, missing_notes, missing_footer = [], [], []
    for index, (slide, page) in enumerate(zip(prs.slides,pdf),1):
        if not slide.has_notes_slide or len(slide.notes_slide.notes_text_frame.text)<100:
            missing_notes.append(index)
        for shape in slide.shapes:
            if min(shape.left,shape.top)<0 or shape.left+shape.width>prs.slide_width+10000 or shape.top+shape.height>prs.slide_height+10000:
                outside.append({"slide":index,"shape":shape.name})
        if "OPTIMIZATION UPDATE" not in page.get_text():
            missing_footer.append(index)
        page.get_pixmap(matrix=fitz.Matrix(1.4,1.4),alpha=False).save(previews/f"slide_{index:02d}.png")
    sheet=Image.new("RGB",(1600,((expected+2)//3)*325),"#D6DEE7")
    draw=ImageDraw.Draw(sheet)
    for index in range(expected):
        image=Image.open(previews/f"slide_{index+1:02d}.png").convert("RGB")
        image.thumbnail((520,292))
        x=8+(index%3)*533;y=22+(index//3)*325
        sheet.paste(image,(x,y))
        draw.text((x+5,y-17),f"Slide {index+1:02d}",fill="#102238")
    sheet.save(output/"contact_sheet.png")
    report={"slides":expected,"pdf_renderer":"LibreOffice25.2.5.2",
            "native_editable_text_tables_diagrams":True,"out_of_bounds_shapes":outside,
            "missing_notes":missing_notes,"missing_footers":missing_footer,
            "automatic_checks_passed":not(outside or missing_notes or missing_footer),
            "visual_contact_sheet":"contact_sheet.png"}
    (output/"qa_report.json").write_text(json.dumps(report,indent=2)+"\n")
    if not report["automatic_checks_passed"]:
        raise AssertionError(report)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output",type=Path,default=ROOT/"results/optimization/presentation")
    parser.add_argument("--no-render",action="store_true",help="Build PPTX/figures only; default also renders PDF and performs QA")
    args=parser.parse_args()
    output=args.output.resolve();output.mkdir(parents=True,exist_ok=True)
    data=snapshot(output)
    if not data["counts"]["train"]:
        raise ValueError("The expanded-data manifest is required for this update deck")
    figures=make_figures(data,output)
    slides=build(data,output,figures)
    (output/"snapshot.json").write_text(json.dumps(data,indent=2,allow_nan=False)+"\n")
    (output/"sources.json").write_text(json.dumps(SOURCES,indent=2)+"\n")
    if not args.no_render:
        render_and_check(output,slides)
    text = ["# Fusion optimization update", "", f"Snapshot: {data['created_utc']}", "",
            "[Editable PowerPoint](fusion_optimization_update.pptx) · [Office-rendered PDF](fusion_optimization_update.pdf) · [Contact sheet](contact_sheet.png)", "",
            "This separate optimization update explains the goal, data expansion, model I/O and architectures, Fourier and internal-inductance refinement, oracle compression, fixed diagnostic-quality gating, matched legacy/fresh results, and official submission status. The original 18-slide first-result deck is preserved.", "",
            "Scores are labeled separately as local development, one-time independent audit, or official public-test results. Oracle diagnostics use validation labels and are not learned results. The original mean/shape blend was selected on legacy validation; fixed diagnostic gates and 75/25 expert blends name both experts explicitly. Sampled fresh, full-native fresh, legacy and audit populations remain separate.", "",
            "Refresh everything after training: `.venv/bin/python scripts/optimization/build_update_presentation.py`.", "",
            "Text, tables, and model diagrams are editable PowerPoint elements. Scientific figures are supplied in `figures/` as PNG/PDF. Speaker notes, source URLs, the complete result snapshot, and automatic QA accompany the deck."]
    (output/"README.md").write_text("\n".join(text)+"\n")
    print(json.dumps({"slides":slides,"output":str(output),"best_legacy":score_text(data.get('best_legacy')),
                      "best_fresh":score_text(data.get('best_fresh')),"latest_full_epoch":max((r.get('epoch',0) for r in data['full_history']),default=0)},indent=2))


if __name__=="__main__":
    main()
