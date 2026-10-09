#!/usr/bin/env python3
"""Build an editable research PowerPoint from the completed, recorded baselines."""
from __future__ import annotations

import json
from pathlib import Path

from PIL import Image, ImageFont
from matplotlib.font_manager import findfont, FontProperties
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "results/presentation"
FIG = OUT / "figures"
NAVY, TEAL, ORANGE = "102238", "12A89D", "E98B42"
GRAY, LIGHT, WHITE, MUTED = "586675", "F0F4F7", "FFFFFF", "D6DEE7"
FONT = "DejaVu Sans"
FONT_FILES = {bold:findfont(FontProperties(family=FONT,weight="bold" if bold else "normal")) for bold in (False,True)}
SOURCES = {
    "challenge": "https://fusion-equilibrium-challenge.sophelio.io/",
    "data": "https://huggingface.co/datasets/Sophelio/fusion-equilibrium-challenge",
    "starter": "https://github.com/Sophelio/fusion-equilibrium-challenge-starter",
    "deep": "https://doi.org/10.1038/s42256-021-00302-5",
    "fourier": "https://doi.org/10.1016/j.cma.2023.116300",
    "codabench": "https://www.codabench.org/competitions/17739/",
}


def rgb(value):
    return RGBColor.from_string(value)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333333)
        self.prs.slide_height = Inches(7.5)
        self.slides = []
        self.notes = []

    def text(self, s, x, y, w, h, value, size=20, color=NAVY, bold=False,
             align=PP_ALIGN.LEFT, middle=False):
        # Fit with the same font as the Office renderer, keeping diagram text inside cards.
        value=str(value)
        requested=size
        while size>max(7,min(12,requested)):
            font=ImageFont.truetype(FONT_FILES[bold],round(size*4))
            n_lines=0
            for paragraph in value.split("\n"):
                line=""
                for word in paragraph.split():
                    candidate=(line+" "+word).strip()
                    if line and font.getlength(candidate)/4>w*72-2:
                        n_lines+=1
                        line=word
                    else:
                        line=candidate
                n_lines+=1
            needed=n_lines*size*1.24+3*(len(value.split("\n"))-1)
            if needed<h*72-1:break
            size-=.5
        box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        if middle:
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, line in enumerate(value.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.alignment = align
            p.line_spacing = 1.0
            p.space_after = Pt(3)
            for run in p.runs:
                run.font.name = FONT
                run.font.size = Pt(size)
                run.font.bold = bold
                run.font.color.rgb = rgb(color)
        return box

    def rect(self, s, x, y, w, h, fill=LIGHT, line=None, rounded=False):
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape._element.spPr.append(OxmlElement("a:effectLst"))
        if line:
            shape.line.color.rgb = rgb(line)
        else:
            shape.line.fill.background()
        if rounded:
            shape.adjustments[0] = 0.08
        return shape

    def arrow(self, s, x1, y1, x2, y2, color=GRAY, width=2.2):
        line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = rgb(color)
        line.line.width = Pt(width)
        tail = OxmlElement("a:tailEnd")
        tail.set("type", "triangle")
        line._element.spPr.get_or_add_ln().append(tail)
        return line

    def card(self, s, x, y, w, h, title, body, accent=TEAL, size=19):
        self.rect(s, x, y, w, h, LIGHT, rounded=True)
        self.rect(s, x, y, .055, h, accent)
        self.text(s, x+.20, y+.15, w-.4, .5, title, 20, accent, True)
        self.text(s, x+.20, y+.77, w-.4, h-.9, body, size)

    def image(self, s, path, x, y, w, h):
        path = Path(path)
        iw, ih = Image.open(path).size
        scale = min(w/iw, h/ih)
        nw, nh = iw*scale, ih*scale
        s.shapes.add_picture(str(path), Inches(x+(w-nw)/2), Inches(y+(h-nh)/2), width=Inches(nw), height=Inches(nh))

    def slide(self, title, kicker="METHODS", source="Local implementation and run artifacts", note=""):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(WHITE)
        self.slides.append(s)
        self.text(s,.55,.18,11.9,.22,kicker,10,TEAL,True)
        self.text(s,.55,.52,12.15,.65,title,29,NAVY,True)
        self.rect(s,.55,1.27,12.2,.025,MUTED)
        self.text(s,.55,6.92,11.65,.20,source,8.3,GRAY)
        self.rect(s,0,7.23,13.34,.27,NAVY)
        self.text(s,.55,7.27,11,.17,"FUSION EQUILIBRIUM CHALLENGE  •  INITIAL DIII-D BASELINES",8,WHITE)
        self.text(s,12.1,7.25,.6,.2,f"{len(self.slides):02d}",9,WHITE,align=PP_ALIGN.RIGHT)
        if note:
            s.notes_slide.notes_text_frame.text = note
        self.notes.append({"slide":len(self.slides),"title":title,"notes":note,"source":source})
        return s

    def table(self, s, x, y, widths, rows, row_h=.52, size=17):
        for ri, row in enumerate(rows):
            xx = x
            for width, value in zip(widths,row):
                self.rect(s,xx,y+ri*row_h,width,row_h,NAVY if ri==0 else (LIGHT if ri%2 else WHITE))
                self.text(s,xx+.12,y+ri*row_h+.10,width-.24,row_h-.13,value,size,
                          WHITE if ri==0 else NAVY,bold=ri==0)
                xx += width

    def save(self):
        OUT.mkdir(parents=True,exist_ok=True)
        self.prs.core_properties.title = "Fusion Equilibrium: DeepONet and Fourier-DeepONet"
        self.prs.core_properties.subject = "Challenge goal, data, model I/O, architectures, and first local results"
        self.prs.core_properties.author = "Fusion project"
        self.prs.core_properties.keywords = "fusion, equilibrium, DeepONet, Fourier-DeepONet, local validation"
        self.prs.save(OUT/"fusion_equilibrium_baselines.pptx")
        (OUT/"slide_content.json").write_text(json.dumps(self.notes,indent=2)+"\n")
        notes = ["# Presentation speaker notes",""]
        for item in self.notes:
            notes += [f"## {item['slide']:02d}. {item['title']}","",item["notes"],"",f"Source: {item['source']}",""]
        (OUT/"speaker_notes.md").write_text("\n".join(notes))
        print(f"Saved {len(self.slides)} slides to {OUT/'fusion_equilibrium_baselines.pptx'}")


def main():
    runs={m:json.loads((ROOT/f"results/{m}/run.json").read_text()) for m in ("deeponet","fourier_deeponet")}
    scores={m:json.loads((ROOT/f"results/{m}/score.json").read_text()) for m in runs}
    metrics={m:json.loads((ROOT/f"results/{m}/validation_predictions.metrics.json").read_text()) for m in runs}
    d=Deck()

    s=d.slide("Fusion equilibrium with neural operators", "RESEARCH PRESENTATION  /  FIRST RESULTS",
              "Source: completed local runs; RESULTS.md; official challenge and starter kit",
              "This presentation explains the task, the released data, the exact model interfaces and implementations, and our completed first results. All reported scores come from 32 held-out DIII-D training shots. These same validation shots selected the checkpoints. We downloaded the entire public release, but trained the first baselines on 256 shots. No Codabench upload or MAST experiment has been performed. The deck's diagrams and text are editable; figures are supplied separately as PNG/PDF/SVG.")
    d.text(s,.7,1.65,7.9,1.5,"From coil currents and laser profiles\nto a 2D magnetic flux map",32,NAVY,True)
    d.text(s,.7,3.35,7.7,.8,"DeepONet + Fourier-DeepONet\nAn initial study on DIII-D",24,GRAY)
    d.rect(s,9.0,1.7,3.55,3.35,NAVY,rounded=True)
    d.text(s,9.2,1.97,3.1,.8,"151",46,WHITE,True,PP_ALIGN.CENTER)
    d.text(s,9.2,2.82,3.1,.35,"input features per frame",16,WHITE,align=PP_ALIGN.CENTER)
    d.text(s,9.2,3.38,3.1,.8,"65 × 65 + 2",27,ORANGE,True,PP_ALIGN.CENTER)
    d.text(s,9.2,4.2,3.1,.45,"flux grid + scalar outputs",15,WHITE,align=PP_ALIGN.CENTER)
    d.card(s,.7,5.1,5.8,1.5,"DeepONet  ·  S = 0.8703","Flux R² 0.9698  |  110,979 parameters",TEAL,17)
    d.card(s,6.75,5.1,5.8,1.5,"Fourier-DeepONet  ·  S = 0.8721","Flux R² 0.9799  |  317,939 parameters",ORANGE,17)

    s=d.slide("Goal: reconstruct magnetic geometry", "01  /  THE CHALLENGE",
              "[1] fusion-equilibrium-challenge.sophelio.io  •  [2] Official dataset card",
              "A tokamak confines plasma using magnetic fields. The challenge asks for an equilibrium reconstruction without the usual magnetic-field diagnostics. We are allowed external coil currents, Thomson-scattering electron temperature and density, plasma current, geometry, and the query timestamps. EFIT is the reference reconstruction that provides training labels. The flux map encodes nested magnetic surfaces. q95 is the edge safety factor and betaN is normalized beta; their inference needs information not contained in flux geometry alone. Challenge 1 is DIII-D reconstruction; Challenge 2 is zero-shot transfer to MAST and requires predictions for both machines. Only Challenge 1 is implemented here. Sources: "+SOURCES['challenge']+" ; "+SOURCES['data'])
    d.card(s,.7,1.75,3.6,2.35,"Allowed measurements","Coil currents + plasma current\nThomson Te and ne profiles\nQuery time and grid",TEAL,19)
    d.arrow(s,4.38,2.88,4.85,2.88)
    d.card(s,4.95,1.75,3.25,2.35,"Learned operator","Infer an equilibrium\nat each EFIT timestamp",NAVY,20)
    d.arrow(s,8.3,2.88,8.77,2.88)
    d.card(s,8.87,1.75,3.75,2.35,"Required predictions","Flux ψ(R, Z): 65 × 65\nSafety factor q95\nNormalized beta βN",ORANGE,19)
    d.text(s,.8,4.55,11.8,.65,"Why it matters: infer equilibrium when magnetic diagnostics are unavailable.",23,NAVY,True)
    d.card(s,.7,5.4,5.8,1.15,"Challenge 1 · this work","DIII-D training → DIII-D test",TEAL,16)
    d.card(s,6.75,5.4,5.8,1.15,"Challenge 2 · future work","DIII-D training → unseen MAST",ORANGE,16)

    s=d.slide("Data: two tokamaks, 9,121 shots, 103.86 GB", "02  /  DATASET",
              "[2] Dataset revision 1e280905…  •  data/download_summary.json",
              "One Parquet row is one plasma shot, with nested arrays for every signal and target map. The full release comprises 7,041 labeled DIII-D training shots, 874 DIII-D inputs-only public-test shots, and 1,206 MAST inputs-only public-test shots. All are downloaded at a pinned revision. Every target frame is diverted, so the irregular EFIT time list need not be contiguous. The primary flux grid is 65×65 for both machines. The public test sets omit EFIT-derived labels, including magnetics_dsep. There is deliberately no MAST training split. Dataset source: "+SOURCES['data'])
    d.image(s,FIG/"dataset_overview.png",.6,1.42,12.0,3.75)
    d.card(s,.7,5.3,3.75,1.25,"One row = one shot","Nested, multirate time series",TEAL,16)
    d.card(s,4.75,5.3,3.75,1.25,"One target frame","65 × 65 grid at an EFIT time",NAVY,16)
    d.card(s,8.8,5.3,3.75,1.25,"Public test labels","Withheld by the organizers",ORANGE,16)

    s=d.slide("The first experiment uses disjoint shots", "03  /  EXPERIMENT DESIGN",
              "data/processed/manifest.json  •  per-run training_manifest.json",
              "The first 288 files in the official lexicographic training order were selected; filenames are anonymized hashes. A seeded whole-shot split gives 256 training shots and 32 validation shots, seed 2026. Training uses up to 64 uniformly selected original frames per shot: 16,213 frames from 55,752 available frames. Validation uses 32 original frames per shot: 1,024 frames from 6,855 available frames. No target interpolation occurs. The validation set is used both for selecting the best epoch and for the local score, so this is not an independent final test. For official test inference, all 874 shots and 193,359 original EFIT timestamps are retained. The remaining labeled shots have not yet been fitted.")
    d.card(s,.7,1.63,3.15,1.45,"7,041 labeled shots","Full DIII-D training release",NAVY,17)
    d.arrow(s,3.96,2.35,4.48,2.35)
    d.card(s,4.58,1.63,3.8,1.45,"First 288 official files","Seeded split by whole shot",NAVY,17)
    d.arrow(s,8.49,2.35,9.0,2.35)
    d.card(s,9.1,1.63,3.45,1.45,"No shot overlap","No train/validation leakage",TEAL,16)
    d.card(s,.7,3.55,3.75,2.08,"TRAIN · 256 shots","16,213 frames\nAt most 64 frames per shot\nFit model and normalization",TEAL,18)
    d.card(s,4.75,3.55,3.75,2.08,"VALIDATE · 32 shots","1,024 frames\n32 frames per shot\nSelect checkpoint + score",ORANGE,18)
    d.card(s,8.8,3.55,3.75,2.08,"TEST · 874 shots","193,359 frames\nEvery original timestamp\nHidden targets",NAVY,18)
    d.text(s,.8,6.0,11.8,.62,"These are subset baselines. The full download does not mean full-corpus training.",21,NAVY,True)

    s=d.slide("Model I/O: features → flux map and two scalars", "04  /  TENSOR CONTRACT",
              "fusion/data.py  •  fusion/models.py  •  scripts/train.py",
              "B denotes the number of frames in a minibatch. Inputs are standardized features of shape B×151 and a shared grid of P=4,225 coordinate pairs (R,Z), scaled to [−1,1] on each axis. The physical DIII-D grid spans R=0.84…2.54 m and Z=−1.6…1.6 m. Rows of every flux array are Z and columns are R. Both neural models return a normalized residual flux B×65×65 and normalized scalars B×2. The training mean map/global flux standard deviation and scalar means/standard deviations restore physical output units. q95 and betaN are separate learned outputs, not derived from the flux map. The scorer derives the boundary and seven consistency quantities from the submitted flux. Per-shot output arrays contain all T EFIT frames.")
    d.card(s,.7,1.6,3.65,2.25,"Branch input","x ∈ ℝ^(B × 151)\nCurrents, profiles, masks, time\nTrain-only standardization",TEAL,19)
    d.card(s,.7,4.2,3.65,1.65,"Trunk coordinates","r ∈ ℝ^(4,225 × 2)\nNormalized (R, Z) locations",NAVY,19)
    d.arrow(s,4.45,2.65,5.05,2.65)
    d.arrow(s,4.45,5.0,5.05,3.8)
    d.rect(s,5.15,2.15,2.8,2.25,NAVY,rounded=True)
    d.text(s,5.35,2.48,2.4,1.5,"DeepONet\nor\nFourier-DeepONet",21,WHITE,True,PP_ALIGN.CENTER)
    d.arrow(s,8.05,2.85,8.67,2.85)
    d.card(s,8.77,1.6,3.78,2.25,"Flux output","ψ ∈ ℝ^(B × 65 × 65)\nRows = Z; columns = R\nUnits: V·s/rad",ORANGE,19)
    d.card(s,8.77,4.2,3.78,1.65,"Scalar output","[q95, βN] ∈ ℝ^(B × 2)\nTwo values per frame",ORANGE,19)
    d.text(s,.85,6.2,11.8,.45,"Restore units:  ψ̂ = μψ(R, Z) + σψ · fθ(x, r)     and     ŝ = μs + σs · hθ(x)",19,NAVY,True)

    s=d.slide("151 features with missing-data handling", "05  /  PREPROCESSING",
              "fusion/data.py: feature_names(), extract_features()  •  starter/data_fixes.py",
              "Feature decomposition: 21 current values and 21 coverage masks; core Te and ne each contribute 16 channel-index bins, 16 masks, and 3 statistics (log1p mean, log1p maximum, coverage), totaling 70; edge Te and ne each contribute 8 bins, 8 masks, and 3 statistics, totaling 38; query time adds one. Total 42+70+38+1=151. The 21 currents comprise 18 F coils, ECOILA, bcoil, and plasma current. Profiles are resampled by channel index rather than full physical chord geometry. Negative Te/ne are treated as invalid measurements; interpolation masks finite samples, sorts/deduplicates timestamps, uses endpoint values outside coverage, and marks those positions unobserved. The official per-shot plasma-current timestamp fix precedes interpolation. Missing channels contribute zero values/masks. Input means and standard deviations are fit on training data; standardized inputs are clipped to ±12. Target-derived dsep and every EFIT label are excluded from features.")
    d.table(s,.75,1.55,[4.3,3.3,4.5],[
        ["Feature group","Dimensions","Construction"],
        ["Current measurements + masks","21 + 21 = 42","18 F coils + E coil + B coil + Ip"],
        ["Core Thomson Te and ne","2 × (16 + 16 + 3) = 70","Bins + masks + summary stats"],
        ["Edge Thomson Te and ne","2 × (8 + 8 + 3) = 38","Bins + masks + summary stats"],
        ["Query timestamp","1","Time in seconds"],
        ["Total","151","Input-only features"]],row_h=.58,size=16)
    d.card(s,.75,5.3,3.78,1.25,"1 · Correct clocks","Apply official per-shot Ip fix",TEAL,16)
    d.card(s,4.77,5.3,3.78,1.25,"2 · Align measurements","Interpolate inputs to EFIT times",NAVY,16)
    d.card(s,8.79,5.3,3.78,1.25,"3 · Normalize","Log1p profiles; train-only scaling",ORANGE,15)

    s=d.slide("DeepONet: coefficients × spatial basis", "06  /  ARCHITECTURE",
              "[3] Lu et al., Nature Machine Intelligence (2021)  •  fusion/models.py",
              "DeepONet has a branch MLP 151→128→128→128→64 with GELU hidden activations and a coordinate trunk 2→128→128→128→64 with Tanh hidden activations. For the shared grid, the branch produces B×64 coefficients and the trunk produces 4,225×64 features. Their inner product is divided by sqrt(64), with a learned scalar bias, then reshaped to B×65×65. A separate branch head 64→128→2 with GELU predicts q95 and betaN in normalized space. This is a learned trunk, not a PCA/fixed POD basis. Training uses a mean-map residual, so physical inference restores the mean map afterward. The implementation has 110,979 parameters. The underlying operator architecture is from Lu et al.: "+SOURCES['deep'])
    d.card(s,.7,1.7,3.55,1.45,"Branch MLP","151 → [128] × 3 → 64",TEAL,16)
    d.card(s,.7,4.3,3.55,1.45,"Coordinate trunk","2 → [128] × 3 → 64",NAVY,16)
    d.text(s,4.6,1.98,2.6,.6,"b(x): B × 64",20,TEAL,True)
    d.text(s,4.6,4.58,2.75,.6,"t(r): 4,225 × 64",19,NAVY,True)
    d.arrow(s,4.3,2.44,4.52,2.44)
    d.arrow(s,4.3,5.0,4.52,5.0)
    d.arrow(s,7.4,2.45,8.1,3.45)
    d.arrow(s,7.4,4.95,8.1,3.7)
    d.card(s,8.22,2.75,4.33,1.9,"Inner product + reshape","fθ(x, r) = Σk bk(x) tk(r) / √64 + c\nNormalized 65 × 65 residual map",ORANGE,17)
    d.text(s,.8,3.42,6.5,.6,"Scalar head:  b(x) → 128 → 2  → [q95, βN]",19,GRAY)
    d.text(s,.85,6.15,11.8,.5,"110,979 parameters  ·  Learned branch and trunk  ·  Shared spatial grid across the batch",19,NAVY,True)

    s=d.slide("Fourier-DeepONet: add a spectral decoder", "07  /  ARCHITECTURE",
              "[4] Zhu et al., CMAME (2023)  •  Spatial adaptation in fusion/models.py",
              "This is a spatial adaptation of Fourier-DeepONet, not the full seismic source-parameter/U-FNO configuration from the paper. It uses the same input branch and coordinate trunk as our vanilla model. Pointwise branch–trunk multiplication gives B×4,225×64. A learned linear lift 64→16 produces B×16×65×65, padded on the upper/right edges to 72×72 for the FFT. Two blocks sum a truncated Fourier convolution and a local 1×1 convolution, then apply GELU. Each spectral layer retains 10 radial real-FFT modes and 10 modes on each positive/negative Z side with learned complex weights. Crop to 65×65 and project with 1×1 convolutions 16→32→1. The scalar head uses the branch directly. Total 317,939 parameters. This is a genuine spectral decoder, not merely Fourier positional encoding. Source: "+SOURCES['fourier'])
    d.card(s,.7,1.65,3.15,1.4,"Branch + trunk","Same encoders as DeepONet",TEAL,17)
    d.arrow(s,3.97,2.35,4.35,2.35)
    d.card(s,4.47,1.65,3.85,1.4,"Pointwise product","b(x) ⊙ t(r): B × 4,225 × 64",NAVY,17)
    d.arrow(s,8.44,2.35,8.82,2.35)
    d.card(s,8.94,1.65,3.6,1.4,"Lift 64 → 16","Reshape: B × 16 × 65 × 65",ORANGE,16)
    d.arrow(s,10.8,3.16,10.8,3.65)
    d.card(s,8.94,3.77,3.6,1.65,"Pad to 72 × 72","Two spectral blocks\n10 Fourier modes per side",ORANGE,17)
    d.arrow(s,8.8,4.56,8.43,4.56)
    d.card(s,4.47,3.77,3.85,1.65,"Crop + project","65 × 65 grid\n1 × 1 convolutions: 16 → 32 → 1",NAVY,16)
    d.arrow(s,4.33,4.56,3.97,4.56)
    d.card(s,.7,3.77,3.15,1.65,"Flux + scalar head","ψ residual map\nBranch → [q95, βN]",TEAL,17)
    d.text(s,.8,5.87,11.7,.8,"Spectral block:  v′ = GELU[ FFT⁻¹(Wk · FFT(v)) + Wlocal v ]\n317,939 parameters  ·  Plain Fourier decoder adapted to spatial coordinates",19,NAVY,True)

    s=d.slide("Training setup and checkpoint selection", "08  /  OPTIMIZATION",
              "scripts/train.py  •  results/{model}/run.json and history.json",
              "Both runs use seed 2026, AdamW with initial learning rate 1e−3 and weight decay 1e−5, cosine decay to 5e−5, gradient-norm clipping at 1.0, and one NVIDIA A40. Target preprocessing subtracts the training mean map and divides by the global training flux standard deviation; q95 and betaN each use training means/std. The objective is normalized flux MSE +0.15 times scalar MSE +1.0 times the sum of horizontal and vertical first-difference MSE. These are differences on the normalized grid, not a Grad–Shafranov PDE residual. All scalar labels in this subset happened to be finite, although the loss supports masking. DeepONet uses batch256 for150epochs and chooses epoch46. Fourier uses batch32 for100epochs and chooses epoch13. Training-loop times are30.97s and376.13s, excluding loading/preparation/scoring/inference. The experiment is not compute-matched.")
    d.table(s,.75,1.55,[4.35,3.8,3.8],[
        ["Setting","DeepONet","Fourier-DeepONet"],
        ["Minibatch / epochs","256 / 150","32 / 100"],
        ["Selected epoch","46","13"],
        ["Parameters","110,979","317,939"],
        ["Training-loop time","31.0 s","376.1 s"],
        ["Hardware / seed","NVIDIA A40 / 2026","NVIDIA A40 / 2026"]],row_h=.55,size=17)
    d.rect(s,.75,5.1,11.95,1.12,NAVY,rounded=True)
    d.text(s,1.0,5.3,11.45,.7,"L = MSEflux + 0.15 · MSEscalars + 1.0 · MSEfirst differences",23,WHITE,True,PP_ALIGN.CENTER)
    d.text(s,.8,6.38,11.8,.4,"Training-only target statistics  ·  Best checkpoint chosen by validation loss  ·  Compute budgets differ",16,GRAY)

    s=d.slide("Scoring the geometry implied by ψ", "09  /  EVALUATION",
              "[1, 2] Official challenge rules  •  starter/fusion_scoring version 3.2.0",
              "The official composite clips the relevant R² contributions at zero and combines four terms:0.55 flux R²,0.15 the mean R² of q95 and betaN,0.10 normalized LCFS agreement,0.20 consistency. LCFS is the last closed flux surface; distance is a normalized Hausdorff discrepancy. The seven consistency quantities are R_axis,Z_axis,elongation kappa,upper/lower triangularity,volume,and internal inductance li. The scorer derives them from both predicted and true flux with the same functionals; they are not separate submitted outputs. It chooses one global flux sign per machine; amplitude is not normalized. We used the unmodified scorer locally on all1,024 cached validationframes. A perfect-prediction check givesS=1 and neither model failed LCFS extraction in this subset. These are not leaderboard scores. Scoring source: "+SOURCES['starter'])
    cards=[("55%","Flux R²","Pooled error over pixels\nand validation frames",TEAL),
           ("15%","q95 + βN","Accuracy of the two\nsubmitted scalar outputs",ORANGE),
           ("10%","LCFS agreement","Boundary distance implied\nby the submitted flux",NAVY),
           ("20%","Consistency","Seven equilibrium quantities\nderived from ψ",GRAY)]
    for i,(weight,title,body,color) in enumerate(cards):
        x=.7+i*3.05
        d.rect(s,x,1.72,2.85,2.65,LIGHT,rounded=True)
        d.text(s,x+.2,1.9,2.45,.75,weight,35,color,True)
        d.text(s,x+.2,2.9,2.45,.55,title,20,NAVY,True)
        d.text(s,x+.2,3.5,2.45,.8,body,15,GRAY)
    d.text(s,.8,4.78,11.85,.68,"S = 0.55 R²ψ + 0.15 R²q95,βN + 0.10 (1 − DLCFS) + 0.20 Consistency",22,NAVY,True)
    d.text(s,.8,5.7,11.85,.88,"Consistency: axis R/Z, elongation, upper/lower triangularity, volume, internal inductance.\nOfficial clipping and sign handling are applied by the supplied scorer.",18,GRAY)

    s=d.slide("First results: a small lead for Fourier-DeepONet", "10  /  LOCAL RESULTS",
              "results/{model}/score.json  •  32 validation shots / 1,024 selected frames",
              "DeepONet scoresS=.870264 with fluxR²=.969776; Fourier scoresS=.872052 with fluxR²=.979898. Fourier improves fluxRMSE from.0330139 to.0269243 V·s/rad, an18.4% reduction. Composite improvement is only.001788 because DeepONet has higher derived-scalar consistency (.542156 vs.518551). q95 R² is.846803 vs.842659; betaN R² is.948950 vs.962249. Neither model has LCFS extraction failures in this validation subset. A training mean-map control has fluxR²=.332368, showing substantial learned predictive signal. The same validation subset selected the checkpoints, and the different compute budgets prevent a controlled claim of architecture superiority.")
    d.card(s,.75,1.5,5.85,1.2,"DeepONet  ·  S = 0.8703","Flux R² 0.9698   |   RMSE 0.0330 V·s/rad",TEAL,17)
    d.card(s,6.85,1.5,5.75,1.2,"Fourier-DeepONet  ·  S = 0.8721","Flux R² 0.9799   |   RMSE 0.0269 V·s/rad",ORANGE,17)
    d.image(s,FIG/"metric_comparison.png",.65,2.92,12.0,3.5)
    d.text(s,.8,6.5,11.85,.3,"Fourier: 18.4% lower flux RMSE, +0.0018 composite S.  Mean-map control: flux R² = 0.3324.",16,NAVY,True)

    s=d.slide("Flux maps on one fixed held-out frame", "11  /  QUALITATIVE RESULTS",
              "Shot diii_d_00019, time 2620 ms, cached frame 17/32  •  shared flux scale",
              "The frame is selected deterministically: the first validation shot and its middle cached frame, not the easiest prediction. The target and model maps share a flux color scale and errors share a symmetric error scale. Read the axes in physical R/Z meters; flux units are V·s/rad. The images show that both learned models recover the main flux structure. Local deviations near the domain edges and shaping-coil regions can be small in global pixel error yet affect derivatives, boundary extraction, and derived quantities. This single frame illustrates behavior but is not the aggregate performance estimate; the previous slide summarizes all1,024 validationframes.")
    d.image(s,FIG/"field_flux.png",.55,1.4,12.22,5.37)

    s=d.slide("Spatial error on the same held-out frame", "11  /  QUALITATIVE RESULTS",
              "Prediction minus target  •  shot diii_d_00019, time 2620 ms  •  units: V·s/rad",
              "This uses the exact same deterministic shot and timestamp as the preceding slide. These are signed errors, prediction minus EFIT target, under the stored DIII-D flux sign. A common symmetric color scale makes the two methods comparable. Spatial derivatives and inferred geometry may be sensitive to errors that contribute little to global field MSE. The full validation metrics, rather than this illustration, determine the reported model comparison.")
    d.image(s,FIG/"field_errors.png",.55,1.4,12.22,5.37)

    s=d.slide("Scalar outputs: safety factor and normalized beta", "12  /  SCALAR PREDICTIONS",
              "All 1,024 validation frames  •  identity line = perfect prediction",
              "These parity plots contain all1,024 cached validationframes for both models, not selected examples. q95 is the edge safety factor at95% flux and betaN is normalized plasma beta. Both are supervised branch-head predictions because a flux map alone cannot provide the toroidal-field function and pressure information needed for them. The identity line indicates perfect predictions. The metrics pool these frames; frames within a shot are correlated, so the scatter does not represent1,024 statistically independent experiments. Finite negative betaN values in the labels were preserved rather than silently removed. Aggregate q95R² slightly favorsDeepONet, whereas betaNR² favorsFourier.")
    d.image(s,FIG/"scalar_parity.png",.6,1.42,12.1,5.28)

    s=d.slide("Derived geometry is the main improvement target", "13  /  ERROR ANALYSIS",
              "Official derived-scalar R² on the same validation set  •  results/{model}/score.json",
              "The consistency term is the mean of nonnegative R² scores for seven functionals extracted from the flux map. DeepONet: axisR .7611, axisZ .5108, elongation .2680, uppertriangularity .7594, lower .6006, volume .4417, li .4535. Fourier: .8038,.3552,.2869,.8223,.6333,.2505,.4779 respectively. The lower Fourier axisZ and volume scores offset its improved flux accuracy. The important observation is that a high global fieldR² does not ensure precise equilibrium-derived quantities. Shape, axis, boundary-aware training or better physical input geometry are possible next experiments, not completed results.")
    d.image(s,FIG/"derived_consistency.png",.6,1.55,12.1,4.65)
    d.text(s,.8,6.25,11.8,.55,"High pixel R² does not guarantee accurate elongation, volume, or magnetic-axis position.",20,NAVY,True)

    s=d.slide("Learning curves favor early checkpoints", "14  /  TRAINING BEHAVIOR",
              "Recorded training histories  •  marked epochs: DeepONet 46, Fourier-DeepONet 13",
              "Training loss continues to fall after the best validation epoch. We retain the checkpoint with the lowest validation objective, not the final epoch. DeepONet's best normalized validation loss is.043263 at epoch46 of150; Fourier's is.033894 at epoch13 of100. Training times shown use the recorded training-and-validation loop only, excluding download,preprocessing,scoring,and test inference. Different batch sizes and epoch counts mean a different number of optimizer steps and a roughly12× training-time difference. This pattern motivates larger and more diverse training subsets, regularization, and an untouched validation split for future model selection.")
    d.image(s,FIG/"training_curves.png",.6,1.55,12.1,4.6)
    d.text(s,.8,6.28,11.8,.45,"Use the best validation checkpoint; later training epochs do not improve these held-out shots.",19,NAVY,True)

    s=d.slide("Completed work and next experiments", "15  /  STATUS AND NEXT STEPS",
              "RESULTS.md  •  submission validation reports  •  proposed next experiments are not completed work",
              "Completed: full pinned dataset download; both trained models and checkpoints; official local metric on32 held-out shots; full DIII-D predictions; independently validated submission ZIPs covering874shots,193,359 timestamps and2,622 float32 arrays per model;10 pipeline checks and perfect-prediction scorer validation. No leaderboard upload has occurred. Limitations: only256 training shots fitted; validation was used for checkpointselection; no repeat-seed uncertainty; compute budgets differ; no MAST inference/score; profile features use channel index rather than full physical geometry. Proposed next steps are to scale training across more of the7,041 labeled shots, reserve an untouched final evaluation split and repeat seeds, incorporate physical coil/chord geometry and shape-aware losses, then investigate zero-shot transfer. Participation upload requires the user's Codabench account. URL: "+SOURCES['codabench'])
    d.card(s,.7,1.62,5.85,2.25,"Completed and reproducible","Both models, checkpoints, plots, and code\nFull 874-shot DIII-D submission ZIPs\n193,359 timestamps; all checks passed",TEAL,19)
    d.card(s,6.8,1.62,5.75,2.25,"Current limits","256 training shots fitted\nValidation also selects checkpoints\nNo leaderboard upload or MAST results",ORANGE,19)
    d.text(s,.82,4.3,11.6,.5,"Next experiments",24,NAVY,True)
    d.table(s,.75,5.0,[.8,5.2,5.95],[
        ["01","Scale training + independent evaluation","More shots, repeated seeds, fixed final split"],
        ["02","Improve geometric consistency","Physical geometry and boundary-aware losses"],
        ["03","Submit and assess transfer","Codabench DIII-D first; MAST zero-shot later"]],row_h=.52,size=15)

    s=d.slide("References and reusable artifacts", "16  /  REPRODUCIBILITY",
              "Dataset CC BY 4.0  •  upstream starter MIT  •  local code, plots, and model outputs retained",
              "Official challenge: "+SOURCES['challenge']+"\nDataset: "+SOURCES['data']+"\nStarter/scorer: "+SOURCES['starter']+"\nDeepONet: Lu et al.2021, "+SOURCES['deep']+"\nFourier-DeepONet: Zhu et al.2023, "+SOURCES['fourier']+"\nDataset revision:1e280905b85f2a6fdde7e06fca8cf3a1edf447cb. Starter commit:a67429165b09eb81c311d44db6ff11743f108b0e. Exact source URLs and all notes are also in speaker_notes.md. The deck summarizes the recorded experiment; no model was retrained to make this presentation.")
    refs=[
        ("[1] Challenge and scoring rules",SOURCES['challenge']),
        ("[2] Dataset + reference implementation",SOURCES['data']+"\n"+SOURCES['starter']),
        ("[3] DeepONet · Lu et al. (2021)",SOURCES['deep']),
        ("[4] Fourier-DeepONet · Zhu et al. (2023)",SOURCES['fourier'])]
    for i,(title,url) in enumerate(refs):
        yy=1.6+i*1.03
        d.text(s,.8,yy,11.85,.35,title,17,NAVY,True)
        box=d.text(s,.8,yy+.4,11.85,.57,url,12,TEAL)
        if "\n" not in url:
            box.text_frame.paragraphs[0].runs[0].hyperlink.address=url
    d.rect(s,.75,5.95,11.85,.75,LIGHT,rounded=True)
    d.text(s,.95,6.12,11.45,.5,"Workspace: RESULTS.md  ·  README.md  ·  results/{model}/  ·  results/presentation/figures/",16,NAVY,True)
    d.save()


if __name__=="__main__":
    main()
